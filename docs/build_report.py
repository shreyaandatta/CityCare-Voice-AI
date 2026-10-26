"""Build docs/AI-SC_Report.docx from the project's code, results and images.
Run:  python docs/build_report.py"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(ROOT / "fuzzy_triage"), str(ROOT / "department_router")]
from fuzzy_triage import MF, RULES, triage_call  # noqa: E402
from samples import SAMPLES  # noqa: E402
from router import load_data, build_model  # noqa: E402
from sklearn.metrics import accuracy_score  # noqa: E402

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text, bold=False, center=False):
    para = doc.add_paragraph()
    run = para.add_run(text); run.bold = bold
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return para


def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def table(header, rows):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
    for c, txt in zip(t.rows[0].cells, header):
        c.text = str(txt); c.paragraphs[0].runs[0].bold = True
    for r in rows:
        for c, txt in zip(t.add_row().cells, r):
            c.text = str(txt)
    doc.add_paragraph()


def image(path, caption, width=6):
    doc.add_picture(str(ROOT / path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(caption, center=True).runs[0].italic = True


# ---------- Title ----------
for _ in range(5):
    doc.add_paragraph()
p("ARTIFICIAL INTELLIGENCE & SOFT COMPUTING", bold=True, center=True).runs[0].font.size = Pt(16)
p("Project Report", center=True)
doc.add_paragraph()
t = p("CityCare Voice AI: An Inbound Hospital Voice Receptionist with Fuzzy-Logic Triage "
      "and Naive Bayes Department Routing", bold=True, center=True)
t.runs[0].font.size = Pt(18)
doc.add_paragraph()
p("Submitted by: Shreyaan Datta", center=True)
p("Source code: github.com/shreyaandatta/CityCare-Voice-AI", center=True)
doc.add_page_break()

# ---------- 1 ----------
h("1. Abstract")
p("Hospitals receive many phone calls that range from routine appointment requests to life-threatening "
  "emergencies. This project builds an AI voice receptionist for a fictional hospital, CityCare. The agent "
  "answers calls in English or Hindi, gives hospital information, and collects patient details. After each "
  "call, a soft-computing pipeline decides how urgent the call is using a Mamdani fuzzy inference system, "
  "routes it to a department using a Multinomial Naive Bayes classifier, suggests a doctor and time slot, "
  "and alerts the emergency desk when needed. Staff see all calls on a dashboard sorted by urgency.")

h("2. Problem Statement and Objectives")
p("Reception desks miss calls at busy times, and urgency is judged inconsistently. Urgency is also vague by "
  "nature: 'a lot of pain' or 'an elderly patient' are not crisp values, which makes fixed if-else thresholds "
  "brittle. The objectives are to:")
bullets(["Answer every inbound call with a conversational voice agent.",
         "Model triage urgency with fuzzy logic so that vague, overlapping symptoms give smooth, explainable decisions.",
         "Route each call to the correct department with a trained machine-learning classifier.",
         "Suggest an appointment slot, flag emergencies, and let callers check their appointment status.",
         "Store every call and show it to staff in priority order."])

h("3. AI & Soft Computing Concepts Used")
table(["Concept", "Use in this project"], [
    ["Intelligent agent", "The voice agent perceives speech and acts by replying and saving records."],
    ["Natural Language Processing", "Speech-to-text, LLM dialogue, text-to-speech (Retell AI)."],
    ["Information extraction", "Post-call analysis extracts name, phone, age, symptoms, request type."],
    ["Fuzzy logic", "Fuzzification, 15-rule Mamdani inference, centroid defuzzification for triage."],
    ["Machine learning", "Multinomial Naive Bayes text classifier for department routing."],
    ["Heuristic knowledge", "Weighted symptom lexicon with negation handling; expert rule base."],
    ["Rule-based reasoning", "Safety override: red-flag symptoms go straight to the Emergency department."],
])
h("3.1 PEAS description of the agent", 2)
table(["Element", "Description"], [
    ["Performance", "Correct triage and department, complete details captured, short calls, no missed emergencies."],
    ["Environment", "Phone calls from patients and relatives; hospital knowledge base; doctor schedules."],
    ["Actuators", "Synthesised voice, webhook to n8n, database records, ER alert email."],
    ["Sensors", "Caller's audio (speech-to-text), caller ID, function-call results."],
])
p("The environment is partially observable, stochastic, sequential and dynamic.")

h("4. System Architecture")
bullets(["Retell AI voice agent 'Riya': speech recognition, LLM receptionist, speech synthesis, post-call analysis.",
         "n8n call pipeline: Webhook → Is call finished? → Extract record → Fuzzy Triage → Department Router → "
         "Load Doctors → Pick Doctor & Slot → Save to CityCare_Calls → Is Emergency? → Alert ER desk.",
         "n8n API workflow: /citycare-status (used by the agent mid-call) and /citycare-calls (dashboard feed, key-protected).",
         "Data tables: CityCare_Calls (call records) and Doctors (department, days, timings).",
         "Staff dashboard: a static web page listing calls by triage score."])

# ---------- 5 ----------
h("5. Fuzzy Triage System")
h("5.1 Feature extraction (crisp inputs)", 2)
p("The caller's words are scored against a weighted lexicon, for example 'chest pain' = 6, 'unconscious' = 8, "
  "'fever' = 4, 'cough' = 1 and 'checkup' = −1. Longer phrases are matched first so 'chest pain' is not also "
  "counted as 'pain', and negated mentions such as 'no fever' are ignored. Age adds to vulnerability "
  "(5 points for under 5 or over 65). Time words such as 'today', 'tomorrow' or 'next week' set the delay.")
h("5.2 Membership functions", 2)
p("All sets are trapezoidal; a triangle is a trapezoid with equal middle points [a, b, b, c].")
table(["Variable", "Fuzzy set", "Points (a, b, c, d)"],
      [[v, term, ", ".join(map(str, pts))] for v, terms in MF.items() for term, pts in terms.items()])
image("fuzzy_triage/images/membership_functions.png", "Figure 1: Membership functions", 5)
h("5.3 Rule base", 2)
table(["#", "IF", "THEN triage is"],
      [[i + 1, " AND ".join(f"{v} is {t}" for v, t in c.items()), o] for i, (c, o) in enumerate(RULES)])
h("5.4 Inference and defuzzification", 2)
p("Mamdani inference is used: AND is the minimum of the memberships, each rule clips its output set, and the "
  "clipped sets are combined with the maximum. The centroid of the combined area gives a crisp score from 0 to 100: "
  "below 40 is Routine, 40 to 70 is Urgent, and above 70 is Emergency.")
image("fuzzy_triage/images/defuzzification_example.png", "Figure 2: Worked defuzzification example", 5)
image("fuzzy_triage/images/control_surface.png", "Figure 3: Control surface (severity × vulnerability)", 5)
h("5.5 Results on sample calls", 2)
rows = []
for name, text, expected in SAMPLES:
    r = triage_call(text)
    rows.append([name, r["severity"], r["distress"], r["vulnerability"], r["delay_days"], r["triage_score"],
                 r["triage_label"], "✓" if r["triage_label"] == expected else "✗"])
table(["Call", "Sev", "Dis", "Vul", "Delay", "Score", "Label", "As expected"], rows)

# ---------- 6 ----------
h("6. Department Router (Naive Bayes)")
Xtr, ytr = load_data("train_data.csv"); Xte, yte = load_data("test_data.csv")
acc = accuracy_score(yte, build_model().fit(Xtr, ytr).predict(Xte))
p("A Multinomial Naive Bayes classifier works on word counts. For each department c it computes "
  "log P(c) + Σ count(w)·log P(w | c) and picks the highest. Laplace-style smoothing (alpha = 0.3) handles "
  "unseen words.")
bullets([f"Training set: {len(Xtr)} sentences generated from symptom lists × everyday sentence templates (English and Hinglish).",
         f"Test set: {len(Xte)} separate hand-written phrases, never used in training.",
         f"Test accuracy: {acc:.1%}.",
         "Emergency calls skip the classifier; if the caller names a department, that is used instead."])
image("department_router/confusion_matrix.png", "Figure 4: Confusion matrix on the unseen test set", 5)

# ---------- 7 ----------
h("7. Implementation")
bullets(["Python (scikit-fuzzy, scikit-learn) for designing, testing and plotting the models.",
         "n8n cloud only runs JavaScript, so a generator script writes the n8n Code nodes from the same lexicon, "
         "membership functions, rules and trained weights. A parity check confirms that the JavaScript and Python "
         "versions give the same labels.",
         "Retell AI single-prompt agent with post-call analysis fields and a custom function, check_appointment_status.",
         "n8n Data Tables for storage; a static HTML dashboard for staff."])

h("8. Testing")
table(["Test", "Result"], [
    ["Fuzzy triage unit tests (labels, negation, phrase priority, age)", "All pass"],
    ["JavaScript vs Python parity (12 calls, 132 router phrases)", "Identical labels and departments"],
    ["Live webhook: routine call (skin rash, next week)", "Routine → Dermatology, slot next week"],
    ["Live webhook: urgent call (4-year-old with high fever, tomorrow)", "Urgent → Paediatrics, slot next day"],
    ["Live webhook: emergency call (70-year-old, chest pain, breathless)", "Emergency → ER, alert branch triggered"],
    ["Status check with known / unknown number", "Returns booking / polite 'not found'"],
    ["Dashboard feed without key", "Rejected"],
])

h("9. Limitations and Future Work")
bullets(["The lexicon and rules are hand-tuned and not clinically validated.",
         "The router's training data is synthetic, so accuracy on real calls would be lower.",
         "Future work: learn the membership functions with ANFIS (neuro-fuzzy) or tune them with a genetic algorithm; "
         "use sentence embeddings instead of keywords; real calendar booking and SMS confirmation."])

h("10. Conclusion")
p("The project combines conversational AI with soft computing. The fuzzy triage system turns vague phone "
  "descriptions into consistent, explainable urgency levels. The Naive Bayes router sends callers to the "
  "right department, and the automated pipeline ensures every call is recorded, prioritised and, for "
  "emergencies, escalated.")

h("References")
bullets(["L. A. Zadeh, 'Fuzzy sets', Information and Control, 1965.",
         "E. H. Mamdani and S. Assilian, 'An experiment in linguistic synthesis with a fuzzy logic controller', 1975.",
         "S. Russell and P. Norvig, Artificial Intelligence: A Modern Approach, 4th ed.",
         "scikit-fuzzy and scikit-learn documentation.", "Retell AI and n8n documentation."])

out = ROOT / "docs" / "AI-SC_Report.docx"
doc.save(out)
print("saved", out)
